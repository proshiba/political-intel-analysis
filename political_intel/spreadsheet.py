from __future__ import annotations

import csv
from pathlib import Path
from urllib.parse import urlparse

from .models import SourceSite

_REQUIRED_COLUMNS = {"url"}
_OPTIONAL_COLUMNS = {"country", "site_name", "category", "notes"}


def load_sites(path: str | Path) -> list[SourceSite]:
    """Load source sites from a CSV or Excel spreadsheet."""
    input_path = Path(path)
    if not input_path.exists():
        raise FileNotFoundError(f"Spreadsheet not found: {input_path}")

    rows = _read_rows(input_path)
    columns = set(rows[0].keys()) if rows else set()
    missing = _REQUIRED_COLUMNS - columns
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")

    sites: list[SourceSite] = []
    for row in rows:
        url = str(row["url"]).strip()
        if not url:
            continue
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise ValueError(f"Invalid URL in spreadsheet: {url}")
        sites.append(
            SourceSite(
                url=url,
                country=str(row.get("country", "") or "Unknown").strip(),
                site_name=str(row.get("site_name", "") or parsed.netloc).strip(),
                category=str(row.get("category", "") or "unspecified").strip(),
                notes=str(row.get("notes", "") or "").strip(),
            )
        )
    return sites


def _read_rows(input_path: Path) -> list[dict[str, str]]:
    if input_path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        from openpyxl import load_workbook

        workbook = load_workbook(input_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip().lower() for value in rows[0]]
        return [
            {headers[index]: str(value or "").strip() for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
        ]

    with input_path.open(newline="", encoding="utf-8-sig") as handle:
        return [{str(key).strip().lower(): str(value or "").strip() for key, value in row.items()} for row in csv.DictReader(handle)]
